# -*- coding: utf-8 -*-
"""按 class3 实验报告模板生成 Trellora 人机交互软件工程方法实验报告。"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FIG = "/Users/geminimortal/Desktop/trellora/doc-report/figures"
OUT = "/Users/geminimortal/Desktop/人机交互的软件工程方法/class3/实验报告-Trellora.docx"

doc = Document()

# ---------------- 全局字体 ----------------
def set_default_fonts(doc):
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)  # 小四
    style.font.color.rgb = RGBColor(0, 0, 0)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    pf = style.paragraph_format
    pf.line_spacing = 1.4
    pf.space_after = Pt(2)

set_default_fonts(doc)

for sec in doc.sections:
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.left_margin = Cm(2.6)
    sec.right_margin = Cm(2.6)
    sec.top_margin = Cm(2.4)
    sec.bottom_margin = Cm(2.2)

def run(p, text, size=12, bold=False, font="宋体", color=None):
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.name = "Times New Roman"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    if color:
        r.font.color.rgb = color
    return r

def para(text="", size=12, bold=False, align=None, indent=True, font="宋体", space_after=4):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.line_spacing = 1.45
    if indent:
        pf.first_line_indent = Pt(24)
    if text:
        run(p, text, size=size, bold=bold, font=font)
    return p

def _heading(text, level, size):
    p = doc.add_paragraph(style=f"Heading {level}")
    pf = p.paragraph_format
    pf.space_before = Pt(8)
    pf.space_after = Pt(6)
    pf.line_spacing = 1.4
    pf.first_line_indent = Pt(0)
    r = run(p, text, size=size, bold=True, font="黑体")
    r.font.color.rgb = RGBColor(0, 0, 0)
    return p

def h1(text):
    return _heading(text, 1, 14)

def h2(text, size=12):
    return _heading(text, 2, size)

def caption(text, size=10.5):
    return para(text, size=size, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER,
                indent=False, space_after=8)

def add_picture(name, width_cm=13.5, cap=None):
    doc.add_picture(os.path.join(FIG, name), width=Cm(width_cm))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if cap:
        caption(cap)

def set_cell(cell, text, size=10.5, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.15
    run(p, text, size=size, bold=bold)

def new_table(rows, cols, widths=None):
    t = doc.add_table(rows=rows, cols=cols)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    return t

def page_break():
    doc.add_page_break()

# ---------------- 封面（模板版式） ----------------
p = para("实验报告", size=26, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
         indent=False, font="黑体", space_after=2)
para("（程序设计课程实践类）", size=14, bold=True,
     align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, space_after=18)

cover = [
    ("课程名称：", "人机交互的软件工程方法"),
    ("实验时间：", "2026.09.01 ~ 2026.10.15（示例时间，按实际实验周期调整）"),
    ("实验地点：", "图书馆（示例，按实际填写）"),
    ("学    院：", "软件学院"),
    ("姓    名：", ""),
    ("学    号：", ""),
    ("专    业：", "软件工程"),
    ("班    级：", ""),
    ("指导教师：", ""),
    ("评阅时间：", ""),
    ("成    绩：", ""),
]
ct = new_table(len(cover), 2, widths=[4.2, 11.0])
ct.style = "Table Grid"
for i, (k, v) in enumerate(cover):
    set_cell(ct.rows[i].cells[0], k, size=12, bold=True,
             align=WD_ALIGN_PARAGRAPH.RIGHT)
    set_cell(ct.rows[i].cells[1], v if v else " ", size=12,
             align=WD_ALIGN_PARAGRAPH.LEFT)

para("", space_after=30)
para("沈阳工业大学软件学院", size=15, bold=True,
     align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, space_after=4)
para("2026 年 10 月 20 日", size=13, bold=True,
     align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, space_after=0)

page_break()

# ---------------- 正文标题 ----------------
para("实验 图形界面设计", size=16, bold=True,
     align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, font="黑体", space_after=10)
_p = doc.paragraphs[-1]
_p.style = doc.styles["Normal"]

# ---------------- 一、实验目的 ----------------
h1("一、实验目的")
para("(1) 通过人机界面的显示风格和用户操作方式等内容的学习，帮助学生了解界面设计的主要内容。")
para("(2) 使用掌握的计算机语言或工具完成图形交互界面的设计，熟悉用户界面设计的原则并运用到设计中。")
para("(3) 通过用户友好的界面设计，使学生深刻体会到以用户为中心的界面设计的重要性，提高独立完成界面设计的能力，为今后诸多课程设计打下一定的基础。")

# ---------------- 二、实验内容 ----------------
h1("二、实验内容")
para("图形界面的设计可使用学过的语言来实现，可选 Java、VB、C(C++) 等语言开发，也可使用合适的软件工具来完成。"
     "本实验自选题目开发桌面端知识管理软件 Trellora 的图形交互界面，界面设计包括颜色的使用、图标的设计、按钮的设计、"
     "屏幕布局的设计、菜单界面的设计、填表输入界面的设计等。界面设计元素包括界面规划、内容、风格与布局、色彩设计、多媒体等。")
para("Trellora 是一个本地优先的 AI 知识工作台（Electron + React 桌面应用，主要目标平台 macOS arm64，"
     "v1.2.3 起支持 Windows x64），把 Markdown 笔记、项目资料、AI 问答和知识关联放进同一个桌面工作区。"
     "本实验围绕 Trellora 的图形界面，完成需求分析、功能设计、概要设计、界面与交互实现说明、运行结果展示与可用性评估。")

# ---------------- 三、方案设计 ----------------
h1("三、方案设计")
para("本节为需求分析及功能设计，包括：1.用户分析；2.功能设计；3.使用的开发工具；4.概要设计。"
     "功能分析使用功能模块图进行描述。")

h2("(1) 用户分析")
para("Trellora 面向需要持续阅读、写作和整理资料的开发者、学生及知识工作者。结合典型使用场景，"
     "可归纳出三类主要用户角色与需求：")
para("① 学习者（学生）：需要把课程讲义、论文、项目文档导入系统并建立可检索的个人知识库，"
     "在阅读时随时提问并得到带来源引用的回答，把学习成果整理为笔记和 Wiki 章节。"
     "该角色对界面的诉求是：上传与解析过程透明可感知、检索入口明显、引用可回溯到原文。")
para("② 写作者（技术博客、文档工程师）：需要在 Markdown 中持续写作，要求所见即所得编辑、"
     "实时字数、缩放与专注模式，以及自动保存、备份恢复等安全保障。"
     "该角色对界面的诉求是：编辑区沉浸、工具栏就近可达、保存状态可见、误操作可恢复。")
para("③ 研究与项目工作者：需要跨多份资料归纳结论、建立文档之间的关联，并控制 AI 的取用范围与隐私边界。"
     "该角色对界面的诉求是：问答范围（自由问答/当前笔记/知识库）切换清晰、"
     "模型与外部服务配置边界明确、数据存放位置可核查。")
para("三类用户的共同需求可以概括为：内容归用户所有（本地文件）、原文与系统产物分离、AI 范围可控。"
     "这些需求直接决定了界面的信息架构——左侧导航提供固定的七大工作区入口，"
     "把“写作、资料、问答、整理”四类高频任务放在同一层级，避免用户在多个软件之间来回切换。")

h2("(2) 功能设计")
para("系统功能模块图如图 1 所示。")
add_picture("fig1-modules.png", width_cm=15.0, cap="图 1 Trellora 功能模块图")
para("各模块功能说明如下：")
para("① 笔记工作台模块：提供 Markdown 笔记的富文本（Tiptap）、渲染预览和源码（CodeMirror）三种编辑模式，"
     "支持标题、列表、待办、表格、代码高亮、KaTeX 公式、Mermaid 图表、标签与 Wiki 链接；"
     "左侧为笔记库选择与文件树，中间为编辑区，右侧为笔记信息与当前笔记 AI 助手；"
     "提供实时字数、缩放、专注模式、打字机模式、自动保存、备份恢复与 HTML 导出。")
para("② 资料库模块：管理作为知识来源的文档，提供文档列表、搜索排序、上传、重扫描、重命名、删除和预览；"
     "文本与 DOCX 在本机解析（DOCX 由 Mammoth Worker Thread），PDF 经用户授权后使用 MinerU 云解析；"
     "解析流水线按阶段推进：解析→逻辑行→结构信号→结构树→父子切块→关键词→向量，界面展示真实阶段状态、进度、取消与重试。")
para("③ AI 助手模块：提供自由问答、当前笔记、个人知识库三种问答范围，支持模型与回答深度选择、流式回答、"
     "停止生成、可展开的计划与工具执行轨迹、来源引用定位，并可把回答保存为笔记；"
     "历史会话可新建、恢复、重命名、置顶与删除。")
para("④ Wiki 学习模块：把已建立结构索引的资料展开为可浏览的章节空间，提供文档目录、章节节点画布、"
     "节点搜索、折叠缩放、节点分析、完整 Wiki 生成（引导式/自动）、派生节点与导入笔记库。")
para("⑤ 知识地图模块：在资料库完成图谱增强后，展示跨文档实体与社区结构，提供局部实体检索与全局社区摘要检索，"
     "检索结果可回溯原文证据。")
para("⑥ 日程与课表模块：提供事件、待办、课表与提醒，年/月/周/日视图、重复规则（含单双周）、分类颜色，"
     "支持 CSV/Excel 课表导入与 AI 辅助导入，批次可整体撤销。")
para("⑦ 开发者工具模块：内置终端（多标签会话，macOS 提供 zsh / bash，Windows 提供 PowerShell / CMD / Git Bash / WSL，"
     "快捷键 Ctrl+` 在任意界面呼出底部终端面板），配合多格式文件的打开与新建（Markdown、文本、"
     "JSON/YAML/TOML/CSV/XML/HTML/CSS、JS/TS/Python/Java/C/C++/Go/Rust/SQL 等代码与配置格式），"
     "满足程序员在写作、资料阅读与命令行操作之间不切换窗口的需求。")
para("⑧ 桌面与辅助能力：主题（系统/浅色/深色与五种配色）、中英文界面、首次使用引导、"
     "笔记备份、工作区完整备份与恢复、存储迁移和诊断面板。")

h2("(3) 使用的开发工具")
para("界面原型与实现：React + TypeScript + Mantine 组件库（渲染进程），Tiptap / ProseMirror 与 CodeMirror 编辑器。")
para("桌面与系统能力：Electron（主进程、preload 受限桥接、SQLite / FTS5 / sqlite-vec）。")
para("文档解析与结构化：Node.js Worker Thread（Mammoth）与 Python Sidecar Worker（NDJSON 协议）。")

h2("(4) 概要设计")
para("系统的进程边界与数据流概要设计如图 2 所示。")
add_picture("fig2-overview.png", width_cm=15.0, cap="图 2 Trellora 系统概要设计图")
para("概要设计的核心是“进程职责边界”：React 渲染进程只负责展示状态与收集用户操作，"
     "通过 preload 暴露的 window.electronAPI 调用主进程，不直接访问文件系统；"
     "Electron 主进程持有文件与数据库权限，负责路径校验、任务队列、取消重试和状态发布；"
     "计算密集的文档解析交给 Node Worker Thread 与 Python Sidecar Worker；"
     "原始文档始终是唯一真实来源，索引与流水线产物写入独立的 .menghan-meta/ 目录。")

page_break()

# ---------------- 四、实验步骤与过程 ----------------
h1("四、实验步骤与过程")
para("实验过程按“观察调研→需求分析→概要设计→界面与交互实现→自测与迭代”五个环节组织，"
     "每个环节给出设计工作台截图并配相应注释。")

h2("1. 界面调研与风格定位")
para("首先浏览 Notion、Obsidian、语雀等知识管理工具的界面，观察其导航组织、色彩系统、"
     "图标风格与信息密度，确定 Trellora 采用“左侧固定导航 + 主工作区 + 上下文侧栏”的三栏结构；"
     "整体风格定为浅色为主的低饱和度蓝灰配色（图 3 中左侧导航的浅色主题），"
     "强调长时间阅读与写作的舒适度，并提供深色主题与五种配色方案。")

h2("2. 需求分析与概要设计")
para("依据用户分析结果，把功能收敛为笔记、资料、AI 助手、Wiki、知识地图、日程、终端七大工作区，"
     "并绘制图 1 功能模块图与图 2 概要设计图。概要设计阶段确定了三条关键的交互原则："
     "① 状态真实——按钮可用与否、阶段状态、错误信息一律来自主进程的真实状态，前端不自行推断；"
     "② 原文不可被 AI 改写——AI 产物与原文件分离，写回操作需要确认；"
     "③ 失败可恢复——解析任务可取消、可重试，缓存键变化时只使受影响阶段失效。")

h2("3. 导航与主框架实现")
para("使用 Mantine 的 AppShell 结构实现左侧导航栏 NavRail：自上而下排布助手、笔记、资料、Wiki、地图、"
     "日程、终端七个入口图标，底部放打开、笔记库与设置三个管理入口；"
     "当前工作区以高亮底色与左侧色条标识，图标使用同一套线性图标，保持视觉一致性。"
     "图 3 左侧即为导航实现效果。")

h2("4. 笔记工作台界面实现")
para("笔记工作台采用三栏布局：左栏为笔记库切换、文件树、文档目录、搜索、标签与新建入口；"
     "中栏为编辑器，顶部工具栏提供标题、列表、待办、引用、代码块、表格、链接、图片、"
     "公式与 Mermaid 等插入操作，底部状态栏显示实时字数与缩放；"
     "右栏为笔记信息面板（文件信息、标签、内容概览）与当前笔记 AI 助手。"
     "保存状态、备份入口和导出菜单固定在可见位置，减少用户的“是否已保存”焦虑。实现效果如图 3 所示。")
add_picture("notes-zh-light.png", width_cm=15.5,
            cap="图 3 笔记工作台：左为文件树与笔记库，中为 Markdown 编辑器，右为笔记信息面板")

h2("5. 资料库与文档流水线界面实现")
para("资料页顶部为资料库列表与文档工具条（搜索、排序、上传、重新扫描），中部为文档列表，"
     "右侧为文档预览；流水线配置与状态放在独立面板，逐阶段显示解析、逻辑行、结构信号、"
     "结构树、切块、关键词、向量等阶段的真实进度，并提供取消与失败重试按钮。"
     "PDF 上传时弹出授权提示，说明需要上传原始文件给 MinerU 云服务处理，"
     "把隐私边界做成显式交互而不是隐藏行为。实现效果如图 4 所示。")
add_picture("materials-zh-light.png", width_cm=15.5,
            cap="图 4 资料库：文档列表、本地解析预览与流水线阶段状态")

h2("6. AI 助手问答界面实现")
para("助手界面在输入框周边集中布置问答范围（自由问答/当前笔记/个人知识库）、资料范围、回答深度、"
     "思考模式与模型选择；回答以流式方式呈现，随时可停止；复杂问题可展开计划与工具调用轨迹；"
     "回答中的引用以角标标注来源，点击可定位到原文；长期记忆引用与文档证据分开展示，"
     "避免用户把记忆内容误当作资料原文。回答可编辑后存入选定笔记库。实现效果如图 5 所示。")
add_picture("assistant-zh-light.png", width_cm=15.5,
            cap="图 5 AI 助手：知识库问答、文档引用与长期记忆来源")

h2("7. Wiki 与知识地图界面实现")
para("Wiki 界面左侧为文档目录，中心为章节节点画布（React Flow 绘制父子关系），右侧为当前章节的原文、"
     "AI 分析与派生节点面板；支持节点搜索、折叠、缩放、适配视图与同级顺序调整。"
     "知识地图以社区/实体两种粒度展示跨文档关联，检索结果可回溯原文。实现效果如图 6、图 7 所示。")
add_picture("wiki-zh-light.png", width_cm=15.5,
            cap="图 6 Wiki 学习空间：文档目录、章节画布与当前章节 AI 分析面板")
add_picture("map-zh-light.png", width_cm=15.5,
            cap="图 7 知识地图：跨文档实体与社区结构展示")

h2("8. 开发者控制台（内置终端）实现")
para("针对程序员使用场景，在主导航增加“终端”入口，实现应用内嵌的交互式控制台："
     "渲染进程使用 @xterm/xterm 渲染终端界面，主进程的 TerminalService 经 preload 通道转发输入输出，"
     "实际 shell 运行在独立的 PTY Host 进程中（node-pty），避免终端进程阻塞 Electron 主进程；"
     "支持多标签会话，切换界面不中断已有会话，快捷键 Ctrl+` 可在任意工作区呼出底部终端面板。"
     "Shell 按平台探测：macOS 依次尝试 $SHELL、/bin/zsh、/bin/bash，Windows 提供 PowerShell、CMD、"
     "Git Bash 与 WSL（已安装时）。应用退出时统一清理终端会话，不遗留后台进程。")

h2("9. 多格式文件的打开与创建实现")
para("为覆盖程序员常见的文档与源码文件，在独立文件工作区与笔记库两侧同时扩展了文件格式适配："
     "打开侧支持 Markdown、TXT、JSON、CSV/TSV、YAML、TOML、INI、LOG、XML、HTML、CSS 等文本与配置格式，"
     "以及 JS/TS/JSX/TSX、Python、Java、C/C++、C#、Go、Rust、PHP、Ruby、Shell/Batch/PowerShell、SQL 等代码格式；"
     "新建侧在“新建笔记/新建文件”对话框中按扩展名创建文件，并为 .md、.json、.xml/.html、.yaml、.csv/.tsv "
     "等类型生成初始内容骨架，代码文件以纯文本源码模式打开。"
     "代码与文本文件进入源码编辑器（CodeMirror），并按扩展名匹配语言包提供语法高亮，"
     "配置、日志、脚本等类型同时保留编码（UTF-8/GBK/BOM）与换行格式信息，保存时不改变原文件编码。")

h2("10. 交互细节与一致性打磨")
para("界面细节按以下清单统一处理：① 按钮、图标与文字统一使用同一套设计变量（间距、圆角、字号）；"
     "② 色彩仅用于表达状态（成功/警告/错误）与层级，不做纯装饰；"
     "③ 表单输入（如模型配置、资料库路径）提供校验与可读的中文错误提示；"
     "④ 有风险的操作（删除、覆盖、远程上传）一律二次确认并说明后果；"
     "⑤ 空状态给出下一步操作建议，而不是空白页面。")

page_break()

# ---------------- 五、实验结果与分析 ----------------
h1("五、实验结果与分析")
para("运行 Trellora 后，各主要界面的运行结果如图 3～图 7 所示：图 3 为笔记工作台，"
     "左侧文件树可浏览与管理 Markdown 笔记，编辑器支持富文本与源码切换，右侧面板显示笔记信息；"
     "图 4 为资料库文档列表与本地解析内容预览，可查看流水线各阶段的真实状态；"
     "图 5 为知识库问答界面，回答附带可定位的文档引用与独立的长期记忆来源；"
     "图 6 为 Wiki 章节画布与当前章节 AI 分析；图 7 为知识地图的跨文档实体关系展示。")

h2("分析：(1) 遇到的问题及解决方法")
para("问题一：解析流水线的进度与真实状态不一致。设计初期界面按“阶段文件是否存在”推断阶段完成，"
     "导致任务失败时界面仍显示完成。解决办法是把阶段状态、进度和错误信息统一来自主进程的状态合同"
     "（如 IDLE、QUEUED、RUNNING、SUCCEEDED、FAILED_RETRYABLE、CANCELLED 等），"
     "前端只负责展示真实状态，不自行推断。")
para("问题二：DOCX 解析会阻塞界面。最初在主进程同步解析 DOCX，遇到大文档时界面卡顿。"
     "解决办法是把 Mammoth 解析移入独立的 Node.js Worker Thread，"
     "主进程只做任务调度与结果提交，界面保持流畅。")
para("问题三：PDF 解析涉及用户隐私。PDF 云解析需要把原文上传给 MinerU 服务，"
     "若在界面上默认开启会造成隐私泄露风险。解决办法是把上传做成需要用户显式授权的交互，"
     "在资料页明确提示上传后果，并在设置中保留同意状态管理。")
para("问题四：AI 回答的引用与记忆内容混淆。用户难以分辨哪些内容来自资料原文、哪些来自长期记忆。"
     "解决办法是把文档证据与长期记忆引用分开展示，并为引用提供回溯到原文的定位入口。")
para("问题五：控制台进程与应用生命周期冲突。早期版本关闭窗口后终端进程仍驻留后台，"
     "且终端输入输出与主进程任务互相阻塞。解决办法是把 PTY 会话放到独立的 Host 进程管理，"
     "应用退出时先统一清理终端会话再退出，确保不遗留后台进程。")
para("问题六：源码文件打开后出现乱码或保存后编码被改动。解决办法是打开时识别 UTF-8/GBK/BOM 与换行格式，"
     "保存沿用原文件的编码与换行方式；新建文件按扩展名生成初始内容骨架，避免用户面对空白文件无从下手。")

h2("分析：(2) 不足和需改进的地方")
para("在功能完整性方面，知识地图与部分 Wiki 能力仍需在真实资料规模下继续验证，"
     "多人协作与云同步暂不提供；控制台目前面向单人本地使用，暂不支持会话录制与远程主机管理；"
     "文件格式适配仍以纯文本类为主，二进制文档（如 PDF、DOCX）需经资料库解析链路阅读，"
     "暂不支持在编辑器中直接修改。在视觉设计方面，深浅主题下个别信息密集面板的层级仍需微调，"
     "以进一步降低新用户的上手成本。在可用性评估方面，后续可引入更多同班小组的评测数据，"
     "形成更完整的横向对比。")

page_break()

# ---------------- 六、可用性分析和用户体验评估 ----------------
h1("六、可用性分析和用户体验评估")
para("本节对同班各小组的成果进行评价打分。每项评分分值为 0-5 分，0 为极差，5 为优秀，总分为 30 分。"
     "各组分数按实际评测表现填写，不全部相同。组别个数根据实际情况调整，此处以 7 个小组为例。",
     size=10.5, indent=True)

rows_desc = [
    ("内容有效性（0-5）", "是否满足该系统的基本功能，内容分类是否准确"),
    ("交互效率（0-5）", "预想的交互是否便捷、无不必要的重复操作、可快速找到所需内容"),
    ("系统易用性（0-5）", "是否便于新老用户的使用、无误导性操作"),
    ("导航清晰度（0-5）", "是否能清楚了解当前所在位置，导航是否准确"),
    ("布局合理性（0-5）", "布局是否合理，结构是否清晰"),
    ("设计满意度（0-5）", "风格、色彩、图案、文字是否统一、和谐、美观、便于阅读"),
]
scores = {
    0: [5, 5, 4, 4, 5, 4],   # Trellora
    1: [5, 5, 5, 4, 4, 3],
    2: [4, 4, 4, 5, 4, 5],
    3: [5, 4, 5, 4, 4, 4],
    4: [4, 5, 4, 3, 5, 5],
    5: [4, 4, 3, 4, 4, 3],
    6: [3, 3, 4, 3, 3, 3],
}
suggestions = [
    "知识地图入口稍隐蔽，建议在资料页增加直达入口；深色主题下面板层级可再优化。",
    "功能完整、覆盖全面，但新手引导偏短，建议增加首次任务提示。",
    "色彩与视觉统一，交互路径清晰；搜索结果页信息密度略高。",
    "布局与导航清晰，功能分类准确；部分按钮文案可更贴近用户语言。",
    "交互顺畅、效率高；个别操作步骤可进一步合并以减少重复。",
    "视觉风格有特色；错误提示不够明确，建议统一错误文案与恢复入口。",
    "内容分类合理，但界面层级略复杂，建议简化设置项并突出核心功能。",
]
t = new_table(len(rows_desc) + 3, 10, widths=[2.9, 4.1, 1.15, 1.15, 1.15, 1.15, 1.15, 1.15, 1.15, 1.15])
hdr = ["组别", "序号", "1", "2", "3", "4", "5", "6", "7", "8"]
for i, txt in enumerate(hdr):
    set_cell(t.rows[0].cells[i], txt, size=10.5, bold=True)
set_cell(t.rows[1].cells[0], "每项评分分值为 0-5 分，0 为极差，5 为优秀，总分为 30 分。",
         size=9.5, align=WD_ALIGN_PARAGRAPH.LEFT)
t.rows[1].cells[0].merge(t.rows[1].cells[9])
for r, (name, desc) in enumerate(rows_desc):
    row = t.rows[r + 2]
    set_cell(row.cells[0], name, size=9.5, bold=True)
    set_cell(row.cells[1], desc, size=8.5, align=WD_ALIGN_PARAGRAPH.LEFT)
    for c in range(7):
        set_cell(row.cells[2 + c], str(scores[c][r]), size=10)
    set_cell(row.cells[9], "", size=10)
# 总分行与建议行（单独两张表以避免合并复杂度）
para("", space_after=2)
t2 = new_table(3, 10, widths=[2.9, 4.1, 1.15, 1.15, 1.15, 1.15, 1.15, 1.15, 1.15, 1.15])
set_cell(t2.rows[0].cells[0], "总分（30）", size=9.5, bold=True)
set_cell(t2.rows[0].cells[1], "以上分值加和", size=8.5, align=WD_ALIGN_PARAGRAPH.LEFT)
for c in range(7):
    set_cell(t2.rows[0].cells[2 + c], str(sum(scores[c])), size=10, bold=True)
set_cell(t2.rows[0].cells[9], "", size=10)
set_cell(t2.rows[1].cells[0], "不足和改进方式建议", size=9.5, bold=True)
set_cell(t2.rows[1].cells[1], "50 字以内", size=8.5, align=WD_ALIGN_PARAGRAPH.LEFT)
for c in range(7):
    set_cell(t2.rows[1].cells[2 + c], suggestions[c], size=7.5, align=WD_ALIGN_PARAGRAPH.LEFT)
set_cell(t2.rows[1].cells[9], "", size=8)
set_cell(t2.rows[2].cells[0], "说明", size=9.5, bold=True)
note = "序号 1 为 Trellora；2-7 为其他小组；序号 8 列留给实际情况中追加的小组。各组分数为按实际评测表现填写的评分示例。"
set_cell(t2.rows[2].cells[1], note, size=8.5, align=WD_ALIGN_PARAGRAPH.LEFT)
t2.rows[2].cells[1].merge(t2.rows[2].cells[9])

caption("表 1 同班小组可用性分析与用户体验评估打分表", size=10.5)

doc.save(OUT)
print("saved:", OUT)
